from io import BytesIO
import unittest
from zipfile import ZipFile

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.text_extractor import TextExtractor


class TextExtractorTests(unittest.TestCase):
    def setUp(self):
        self.extractor = TextExtractor()

    def pdf_bytes(self, texts, encrypted=False):
        writer = PdfWriter()
        for text in texts:
            page = writer.add_blank_page(width=612, height=792)
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
            })
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = stream
        if encrypted:
            writer.encrypt("password")
        result = BytesIO()
        writer.write(result)
        return result.getvalue()

    def docx_bytes(self, document):
        result = BytesIO()
        document.save(result)
        return result.getvalue()

    def test_pdf_all_pages(self):
        text = self.extractor.extract(self.pdf_bytes(["Python developer", "Experience"]))
        self.assertEqual(text, "Python developer\nExperience")

    def test_docx_paragraphs_and_tables_in_order(self):
        document = Document()
        document.add_paragraph("Разработчик")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Python"
        table.cell(0, 1).text = "PostgreSQL"
        document.add_paragraph("Опыт работы")
        self.assertEqual(
            self.extractor.extract(self.docx_bytes(document)),
            "Разработчик\nPython\nPostgreSQL\nОпыт работы",
        )

    def test_empty_text_rejected(self):
        document = Document()
        document.add_paragraph(" \t\n ")
        for data in (self.pdf_bytes([""]), self.docx_bytes(Document()), self.docx_bytes(document)):
            with self.subTest(data=data[:10]):
                with self.assertRaisesRegex(ValueError, "Resume text must not be empty"):
                    self.extractor.extract(data)

    def test_invalid_files_rejected(self):
        archive = BytesIO()
        with ZipFile(archive, "w") as target:
            target.writestr("test.txt", "text")
        for data in (b"", b"plain text", b"%PDF-1.7\nbroken", archive.getvalue(), self.pdf_bytes(["text"], encrypted=True)):
            with self.subTest(data=data[:10]):
                with self.assertRaises(ValueError):
                    self.extractor.extract(data)


if __name__ == "__main__":
    unittest.main()
